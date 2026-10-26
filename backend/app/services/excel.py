"""추천 결과 엑셀(.xlsx) 생성 - pandas → openpyxl, 메모리 버퍼."""
from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


def recommendations_to_xlsx(items: list[dict], user_name: str) -> bytes:
    df = pd.DataFrame([{
        "순위": i + 1,
        "상품명": r["title"],
        "여행사": r["agency"],
        "지역": r["region"],
        "도시": r["city"],
        "출발일": f'{r["depart_date"]} ~ {r["depart_end"]}' if r.get("depart_end") else r["depart_date"],
        "일정": f"{r['nights']}박 {r['days']}일" if r.get("nights") is not None else "",
        "판매가(원)": r["price"],
        "할인율(%)": r["discount_rate"],
        "추천점수": r["score"],
        "추천사유": ", ".join(r["reasons"]),
        "링크": r["detail_url"] or "",
    } for i, r in enumerate(items)])

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="맞춤추천", startrow=2)
        ws = writer.sheets["맞춤추천"]
        ws["A1"] = f"NomadPick 맞춤 추천 Top {len(items)} - {user_name}"
        ws["A1"].font = Font(bold=True, size=14)
        header_fill = PatternFill("solid", fgColor="0D6EFD")
        for cell in ws[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        widths = {"A": 6, "B": 42, "C": 10, "D": 8, "E": 10, "F": 12, "G": 10, "H": 13, "I": 10, "J": 10, "K": 48, "L": 40}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
            row[7].number_format = "#,##0"
            link = row[11]
            if link.value:
                link.hyperlink = link.value
                link.style = "Hyperlink"
    return buf.getvalue()
