from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient

from backend.app.main import app
from backend.app.utils import now_kst
from rpa.run import run_deals, run_public, run_report
from rpa.sink import DirectSink

RPA = {"X-RPA-KEY": "test-rpa-key"}


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        sink = DirectSink()
        now = now_kst()
        run_deals(sink, mode="sample", at=now - timedelta(days=1))
        run_deals(sink, mode="sample", at=now)
        run_public(sink, mode="sample")
        run_report(sink)
        yield c


def login(client, email="demo@test.com", password="Demo!2026"):
    r = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def test_health_and_static(client):
    assert client.get("/api/v1/health").json()["db_ok"] is True
    r = client.get("/")
    assert r.status_code == 200 and "NomadPick" in r.text


def test_signup_login_flow(client):
    r = client.post("/api/v1/auth/signup", json={"email": "New@Test.com", "password": "password1", "name": "신규"})
    assert r.status_code == 201 and r.json()["has_preference"] is False
    assert client.post("/api/v1/auth/signup", json={"email": "new@test.com", "password": "password1", "name": "x"}).status_code == 409
    assert client.post("/api/v1/auth/login", json={"email": "new@test.com", "password": "wrong"}).json()["code"] == "INVALID_CREDENTIALS"
    h = login(client, "new@test.com", "password1")
    assert client.get("/api/v1/users/me", headers=h).json()["email"] == "new@test.com"
    assert client.get("/api/v1/users/me").status_code == 401
    assert client.get("/api/v1/users/me", headers={"Authorization": "Bearer broken"}).json()["code"] == "INVALID_TOKEN"


def test_preferences_save_and_affect_recommendations(client):
    h = login(client, "new@test.com", "password1")
    assert client.get("/api/v1/preferences", headers=h).json() is None
    no_pref = client.get("/api/v1/recommendations", headers=h).json()
    assert "선호 미설정" in no_pref[0]["reasons"][-1]

    body = {"regions": ["유럽"], "budget_min": 4000000, "budget_max": 1000000, "depart_months": [], "nights_min": 6,
            "nights_max": 9, "theme": "관광", "companion": "가족"}
    saved = client.put("/api/v1/preferences", json=body, headers=h).json()
    assert (saved["budget_min"], saved["budget_max"]) == (1000000, 4000000)  # 역전 입력 보정
    recs = client.get("/api/v1/recommendations", headers=h).json()
    assert recs and all(r["region"] == "유럽" for r in recs[:3])
    assert client.put("/api/v1/preferences", json={**body, "depart_months": ["2026/12"]}, headers=h).status_code == 422


def test_recommendations_rules(client):
    recs = client.get("/api/v1/recommendations", headers=login(client)).json()
    assert len(recs) == 10
    assert recs == sorted(recs, key=lambda r: r["score"], reverse=True)
    assert all(r["seats_left"] != 0 for r in recs)
    keys = [(r["agency_code"], r["product_code"]) for r in recs]
    assert len(keys) == len(set(keys))
    cities = [r["city"] for r in recs]
    assert max(cities.count(c) for c in cities) <= 2
    assert any("선호 지역" in x for x in recs[0]["reasons"])


def test_recommendation_excel(client):
    r = client.get("/api/v1/recommendations/excel", headers=login(client))
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    import io

    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 13 and ws["B3"].value == "상품명"


def test_products_filters_sort_and_paging(client):
    data = client.get("/api/v1/products", params={"region": "동남아,일본", "price_max": 1000000, "size": 5}).json()
    assert data["total"] > 0 and len(data["items"]) == 5
    assert all(i["region"] in ("동남아", "일본") and i["price"] <= 1000000 for i in data["items"])
    rates = [i["discount_rate"] for i in data["items"]]
    assert rates == sorted(rates, reverse=True)
    by_price = client.get("/api/v1/products", params={"sort": "price", "agency": "modetour", "size": 50}).json()["items"]
    assert [i["price"] for i in by_price] == sorted(i["price"] for i in by_price) and {i["agency"] for i in by_price} == {"모두투어"}
    month = by_price[0]["depart_date"][:7]
    assert all(i["depart_date"].startswith(month) for i in client.get("/api/v1/products", params={"month": month}).json()["items"])
    assert client.get("/api/v1/products", params={"q": "다낭"}).json()["items"][0]["city"] == "다낭"
    assert client.get("/api/v1/products", params={"month": "2026-13x"}).status_code == 422
    meta = client.get("/api/v1/meta").json()
    assert meta["total"] > 150 and len(meta["agencies"]) == 3


def test_product_detail_prices_and_flights(client):
    item = client.get("/api/v1/products", params={"region": "국내", "sort": "depart", "size": 1}).json()["items"][0]
    d = client.get(f"/api/v1/products/{item['id']}").json()
    assert d["is_domestic"] and d["flights"] and d["flights"][0]["dep_airport_nm"] == "김포"
    prices = client.get(f"/api/v1/products/{item['id']}/prices").json()["items"]
    assert len(prices) == 2
    assert client.get("/api/v1/products/999999").json()["code"] == "PRODUCT_NOT_FOUND"


def test_favorites_and_price_drop(client):
    h = login(client)
    pid = client.get("/api/v1/products", params={"size": 1}).json()["items"][0]["id"]
    assert client.post("/api/v1/favorites", json={"product_id": pid}, headers=h).status_code == 201
    client.post("/api/v1/favorites", json={"product_id": pid}, headers=h)
    favs = client.get("/api/v1/favorites", headers=h).json()
    assert len(favs) == 1 and favs[0]["price_dropped"] is False
    assert client.get(f"/api/v1/products/{pid}", headers=h).json()["is_favorite"] is True
    assert client.delete(f"/api/v1/favorites/{pid}", headers=h).status_code == 204
    assert client.get("/api/v1/favorites", headers=h).json() == []


def test_rpa_api_requires_key_and_upserts(client):
    body = {"agency_code": "hanatour", "collected_at": datetime.now().isoformat(), "source": "live", "items": [{
        "product_code": "HT-TEST01", "title": "테스트 다낭 3박5일", "region": "동남아", "city": "다낭", "theme": "휴양",
        "depart_date": (now_kst() + timedelta(days=30)).date().isoformat(), "nights": 3, "days": 5, "price": 500000, "original_price": 1000000}]}
    assert client.post("/api/v1/rpa/products/bulk", json=body).status_code == 401
    assert client.post("/api/v1/rpa/products/bulk", json=body, headers={"X-RPA-KEY": "wrong"}).status_code == 401
    r1 = client.post("/api/v1/rpa/products/bulk", json=body, headers=RPA).json()
    assert r1["inserted"] == 1 and r1["snapshots"] == 1
    body["items"][0]["price"] = 450000
    r2 = client.post("/api/v1/rpa/products/bulk", json=body, headers=RPA).json()
    assert r2 == {**r2, "inserted": 0, "updated": 1}
    found = client.get("/api/v1/products", params={"q": "테스트 다낭"}).json()["items"][0]
    assert found["price"] == 450000 and found["discount_rate"] == 55.0
    assert client.post("/api/v1/rpa/products/bulk", json={**body, "agency_code": "nope"}, headers=RPA).status_code == 400
    log = {"job_name": "BOT-01", "agency_code": "hanatour", "status": "SUCCESS", "collected": 1, "started_at": datetime.now().isoformat()}
    assert client.post("/api/v1/rpa/logs", json=log, headers=RPA).status_code == 201


def test_anytime_domestic_product_keeps_one_row(client):
    """출발일 표기 없는 국내 상품(상시 출발)은 출발일이 매일 바뀌어도 같은 상품 행·가격 이력을 이어 쓴다."""
    today = now_kst()
    item = {"product_code": "HT-ANY01", "title": "상시 홍도/흑산도 3일", "region": "국내", "city": "홍도", "price": 364000,
            "image_url": "https://image.hanatour.com/a.jpg", "depart_anytime": True}

    def post(at, price):
        body = {"agency_code": "hanatour", "collected_at": at.isoformat(), "source": "live", "items": [{
            **item, "price": price, "depart_date": at.date().isoformat(), "depart_end": (at + timedelta(days=60)).date().isoformat()}]}
        return client.post("/api/v1/rpa/products/bulk", json=body, headers=RPA).json()

    assert post(today - timedelta(days=1), 364000)["inserted"] == 1
    r = post(today, 355000)
    assert (r["inserted"], r["updated"]) == (0, 1)
    found = client.get("/api/v1/products", params={"q": "상시 홍도"}).json()["items"]
    assert len(found) == 1 and found[0]["depart_date"] == today.date().isoformat() and found[0]["price"] == 355000
    assert found[0]["image_url"] == "https://image.hanatour.com/a.jpg"
    assert len(client.get(f"/api/v1/products/{found[0]['id']}/prices").json()["items"]) == 2


def test_admin_endpoints(client):
    assert client.get("/api/v1/admin/rpa-logs", headers=login(client)).status_code == 403
    h = login(client, "admin@test.com", "Admin!2026")
    d = client.get("/api/v1/admin/rpa-logs", headers=h).json()
    assert d["summary"]["today_runs"] >= 3 and len(d["agencies"]) == 3
    assert {l["job_name"] for l in d["logs"]} >= {"BOT-01", "BOT-02", "BOT-03"}
    reports = client.get("/api/v1/admin/reports", headers=h).json()
    day = reports[0]["report_date"]
    xlsx = client.get(f"/api/v1/admin/reports/{day}", headers=h)
    assert xlsx.status_code == 200 and len(xlsx.content) > 10000
    assert "plotly" in client.get(f"/api/v1/admin/reports/{day}", params={"format": "html"}, headers=h).text.lower()
    assert client.get("/api/v1/admin/reports/2000-01-01", headers=h).status_code == 404


def test_auto_mode_keeps_live_data_when_live_fails(client):
    """auto 모드에서 live 수집이 실패하면 기존 live 상품을 sample 로 덮어쓰지 않는다."""
    from rpa.run import _run_deals

    body = {"agency_code": "ybtour", "collected_at": datetime.now().isoformat(), "source": "live", "items": [{
        "product_code": "YB-LIVE01", "title": "라이브 다낭 3박5일", "region": "동남아", "city": "다낭",
        "depart_date": (now_kst() + timedelta(days=20)).date().isoformat(), "nights": 3, "days": 5, "price": 599000}]}
    assert client.post("/api/v1/rpa/products/bulk", json=body, headers=RPA).status_code == 200
    out = _run_deals(DirectSink(), "auto", now_kst(), True, None)  # 브라우저 없음 → live 실패
    assert out["ybtour"]["source"] == "live-kept"
    found = client.get("/api/v1/products", params={"q": "라이브 다낭"}).json()["items"]
    assert found and found[0]["source"] == "live"


def test_isolated_rpa_job_is_killed_on_timeout(monkeypatch):
    """RPA 별도 프로세스가 시간 한도를 넘으면 강제 종료되고 FAIL 로그가 남는다 (웹 서버 보호)."""
    from backend.app.config import settings
    from rpa import run as rpa_run

    monkeypatch.setattr(settings, "rpa_isolated", True)
    monkeypatch.setattr(settings, "rpa_job_timeout_sec", 0)
    out = rpa_run.run_job_isolated("report")
    assert "시간 초과" in out["killed"]
    assert not rpa_run._lock.locked()
