"""산출물 생성 - 진행일지 4종·산출물 4종(hwpx), 발표자료(pptx) 내용 수정 후 PDF 변환.

    .venv\\Scripts\\python.exe docs\\tools\\build_docs.py            # 전체
    .venv\\Scripts\\python.exe docs\\tools\\build_docs.py --no-pdf   # PDF 변환 생략

샘플 원본은 Git 최초 커밋(HEAD 의 첫 커밋) 파일을 사용하므로 여러 번 실행해도 결과가 같다.
"""
import argparse
import copy
import re
import shutil
import subprocess
import sys
import tempfile
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))

import content_spec as spec  # noqa: E402
import diagrams  # noqa: E402
from hwpx_edit import HwpxDoc  # noqa: E402
from mermaid_render import render  # noqa: E402

ROOT = TOOLS.parents[1]
DOCS = ROOT / "docs"
SCREENS = DOCS / "md" / "images"
LEFTOVER = re.compile(r"트래블핏|TravelFit|travelfit|여행경보|환율|외교부|수출입|cron")
DIAGRAM_SIZE = {"SERVICE_FLOW": (1800, 700), "GANTT": (1400, 700), "ARCHITECTURE": (1600, 1300),
                "SEQUENCE": (1900, 1000), "ERD": (1800, 1400)}


def first_commit() -> str:
    return subprocess.run(["git", "rev-list", "--max-parents=0", "HEAD"], cwd=ROOT, capture_output=True, text=True,
                          check=True).stdout.split()[0]


def template(rel: str) -> Path:
    """원본 샘플 파일을 임시 경로로 꺼낸다."""
    data = subprocess.run(["git", "show", f"{first_commit()}:{rel}"], cwd=ROOT, capture_output=True, check=True).stdout
    out = Path(tempfile.gettempdir()) / ("np_tpl_" + Path(rel).name)
    out.write_bytes(data)
    return out


_diagram_cache: dict[str, bytes] = {}


def diagram(name: str) -> bytes:
    if name not in _diagram_cache:
        w, h = DIAGRAM_SIZE[name]
        _diagram_cache[name] = render(getattr(diagrams, name), w, h)
        (TOOLS / "out").mkdir(exist_ok=True)
        (TOOLS / "out" / f"{name.lower()}.png").write_bytes(_diagram_cache[name])
    return _diagram_cache[name]


def build_hwpx(rel: str, conf: dict) -> Path:
    doc = HwpxDoc(template(rel))
    cells, paras = conf.get("cells", {}), {**spec.GLOBAL_PARAS, **conf.get("paras", {})}
    doc.apply_cells(cells)
    doc.apply_paras(paras)
    doc.apply_subs(spec.GLOBAL_SUBS)
    for item_id, name in conf.get("images", {}).items():
        doc.replace_image(item_id, diagram(name))
    missing = [k for k in doc.unused(cells, conf.get("paras", {}))]
    if missing:
        print(f"  [경고] {rel} 미적용 키 {len(missing)}건:", *missing, sep="\n    - ")
    left = sorted({m for m in LEFTOVER.findall(doc.all_text())})
    if left:
        print(f"  [확인] {rel} 잔여 용어: {left}")
    out = ROOT / rel
    doc.save(out)
    print("HWPX:", out)
    return out


# ------------------------------------------------------------------ PPTX ----
def _replace_runs(paragraph, subs):
    for run in paragraph.runs:
        new = run.text
        for old, rep in subs:
            new = new.replace(old, rep)
        if new != run.text:
            run.text = new


def _replace_picture(slide, pic, png: bytes):
    w_px, h_px = Image.open(BytesIO(png)).size
    box_w, box_h = pic.width, pic.height
    scale = min(box_w / w_px, box_h / h_px)
    w, h = int(w_px * scale), int(h_px * scale)
    left = pic.left + (box_w - w) // 2 if pic.left > 700000 else pic.left
    new = slide.shapes.add_picture(BytesIO(png), Emu(left), Emu(pic.top), Emu(w), Emu(h))
    pic._element.addprevious(new._element)
    pic._element.getparent().remove(pic._element)


def _screens_slide(prs, template_slide, position: int):
    shots = [("오늘의 특가 (SC-04)", "mockup_index.png"), ("맞춤 추천 Top 10 (SC-05)", "mockup_recommend.png"),
             ("상품 상세 · 가격 추이 (SC-06)", "mockup_product.png"), ("관리자 수집 모니터링 (SC-08)", "mockup_admin.png")]
    if not all((SCREENS / f).is_file() for _, f in shots):
        print("  [건너뜀] 화면 캡처 없음:", SCREENS)
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for shp in template_slide.shapes:  # 상단 바·제목·구분선·푸터·쪽번호 복제
        if shp.name in {"Rectangle 1", "Rectangle 2", "TextBox 3", "TextBox 4", "Connector 5", "TextBox 6", "TextBox 7"}:
            slide.shapes._spTree.append(copy.deepcopy(shp._element))
    for shp in slide.shapes:
        if shp.name == "TextBox 3":
            shp.text_frame.paragraphs[0].runs[0].text = "반응형 주요 화면"
        elif shp.name == "TextBox 4":
            shp.text_frame.paragraphs[0].runs[0].text = "PC · 모바일 동일 서비스"
        elif shp.name == "TextBox 7":
            shp.text_frame.paragraphs[0].runs[0].text = str(position + 1)

    x0, y0, gap = 548640, 1325880, 182880
    cell_w = (11094415 - gap) // 2
    cell_h = (6309360 - y0 - gap) // 2
    cap_h = 320040
    for i, (caption, file) in enumerate(shots):
        x = x0 + (i % 2) * (cell_w + gap)
        y = y0 + (i // 2) * (cell_h + gap)
        img = Image.open(SCREENS / file)  # 노트북+스마트폰 목업(투명 배경) - 비율 유지
        box_h = cell_h - cap_h
        w = min(cell_w, int(box_h * img.width / img.height))
        h = int(w * img.height / img.width)
        slide.shapes.add_picture(str(SCREENS / file), Emu(x + (cell_w - w) // 2), Emu(y + (box_h - h) // 2), Emu(w), Emu(h))
        tb = slide.shapes.add_textbox(Emu(x), Emu(y + cell_h - cap_h), Emu(cell_w), Emu(cap_h))
        tb.text_frame.paragraphs[0].alignment = 2
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = caption
        run.font.size = Pt(12)
        run.font.bold = True
    ids = prs.slides._sldIdLst
    el = ids[-1]
    ids.remove(el)
    ids.insert(position, el)


def build_pptx() -> Path:
    rel = "docs/프로젝트_발표자료.pptx"
    prs = Presentation(template(rel))
    for slide in prs.slides:
        for shp in slide.shapes:
            if shp.has_text_frame:
                for p in shp.text_frame.paragraphs:
                    if p.text in spec.PPTX_TEXT_EXACT and p.runs:
                        p.runs[0].text = spec.PPTX_TEXT_EXACT[p.text]
                        for r in p.runs[1:]:
                            r.text = ""
                    _replace_runs(p, spec.PPTX_SUBS)
            if getattr(shp, "has_table", False) and shp.has_table:
                for row in shp.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            if p.text in spec.PPTX_CELL_EXACT and p.runs:
                                p.runs[0].text = spec.PPTX_CELL_EXACT[p.text]
                            _replace_runs(p, spec.PPTX_SUBS)

    slides = list(prs.slides)
    for idx, name in ((5, "ARCHITECTURE"), (7, "SEQUENCE")):
        pic = next(s for s in slides[idx].shapes if s.shape_type == 13)
        _replace_picture(slides[idx], pic, diagram(name))

    # 시연(12) 뒤에 '주요 화면' 슬라이드 추가 → 이후 쪽번호 +1
    _screens_slide(prs, slides[11], 12)
    for slide in list(prs.slides)[13:]:
        for shp in slide.shapes:
            if shp.name == "TextBox 7" and shp.text_frame.text.isdigit():
                shp.text_frame.paragraphs[0].runs[0].text = str(int(shp.text_frame.text) + 1)

    for slide in prs.slides:
        for shp in slide.shapes:
            if shp.has_text_frame:
                for p in shp.text_frame.paragraphs:
                    if p.text in spec.PPTX_FINAL and p.runs:
                        p.runs[0].text = spec.PPTX_FINAL[p.text]
                        for r in p.runs[1:]:
                            r.text = ""

    text = "\n".join(shp.text_frame.text for s in prs.slides for shp in s.shapes if shp.has_text_frame)
    text += "\n".join(c.text for s in prs.slides for shp in s.shapes if getattr(shp, "has_table", False) and shp.has_table
                      for r in shp.table.rows for c in r.cells)
    left = sorted(set(LEFTOVER.findall(text)))
    if left:
        print(f"  [확인] pptx 잔여 용어: {left}")
    out = ROOT / rel
    prs.save(out)
    print("PPTX:", out)
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()

    hwpx_out = [build_hwpx(f"docs/{name}", conf) for name, conf in spec.JOURNAL.items()]
    hwpx_out += [build_hwpx(f"docs/산출물/{name}", conf) for name, conf in spec.DELIVERABLES.items()]
    pptx_out = build_pptx()
    if args.no_pdf:
        return

    from office_com import hwpx_refresh_and_pdf, pptx_animate, pptx_to_pdf

    pptx_animate(pptx_out)
    shutil.copyfile(pptx_out, DOCS / "presentation" / "nomadpick.pptx")  # GitHub Pages 발표용(영문 파일명)

    pairs = []
    for path in hwpx_out:
        rel = path.relative_to(DOCS)
        pairs.append((path, DOCS / "pdf" / rel.with_suffix(".pdf")))
    hwpx_refresh_and_pdf(pairs)
    pptx_to_pdf(pptx_out, DOCS / "pdf" / "프로젝트_발표자료.pdf")


if __name__ == "__main__":
    main()
