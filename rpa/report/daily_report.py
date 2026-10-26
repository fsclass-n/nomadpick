"""BOT-03 일일 특가 리포트 - pandas 집계 + Seaborn·Matplotlib 차트(엑셀 삽입) + Plotly HTML."""
from datetime import date
from io import BytesIO

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
import plotly.express as px  # noqa: E402
import seaborn as sns  # noqa: E402
from matplotlib import font_manager, ticker  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

KOREAN_FONTS = ["Malgun Gothic", "NanumGothic", "AppleGothic", "Noto Sans CJK KR", "Noto Sans KR"]


def _setup_font() -> None:
    installed = {f.name for f in font_manager.fontManager.ttflist}
    for name in KOREAN_FONTS:
        if name in installed:
            plt.rcParams["font.family"] = name
            break
    plt.rcParams["axes.unicode_minus"] = False


def _png(fig) -> BytesIO:
    buf = BytesIO()
    fig.savefig(buf, format="png", dpi=110, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def build_report(products: list[dict], report_date: date) -> tuple[bytes, str, dict]:
    """(xlsx 바이트, plotly html, 요약) 반환."""
    _setup_font()
    df = pd.DataFrame(products)
    if df.empty:
        df = pd.DataFrame(columns=["agency", "region", "city", "title", "price", "discount_rate", "depart_date", "seats_left", "detail_url"])
    df["price"] = pd.to_numeric(df["price"], errors="coerce")
    df["discount_rate"] = pd.to_numeric(df["discount_rate"], errors="coerce").fillna(0)

    summary = {
        "report_date": report_date.isoformat(),
        "total": int(len(df)),
        "agencies": int(df["agency"].nunique()),
        "avg_price": int(df["price"].mean()) if len(df) else 0,
        "avg_discount": round(float(df["discount_rate"].mean()), 1) if len(df) else 0.0,
        "max_discount": round(float(df["discount_rate"].max()), 1) if len(df) else 0.0,
    }
    by_agency = (df.groupby("agency").agg(상품수=("title", "count"), 평균가=("price", "mean"), 평균할인율=("discount_rate", "mean"),
                                          최대할인율=("discount_rate", "max")).round(1).reset_index().rename(columns={"agency": "여행사"}))
    by_region = (df.groupby("region").agg(상품수=("title", "count"), 최저가=("price", "min"), 평균가=("price", "mean"),
                                          평균할인율=("discount_rate", "mean")).round(1).reset_index().rename(columns={"region": "지역"})
                 .sort_values("상품수", ascending=False))
    top = df.sort_values("discount_rate", ascending=False).head(20)[
        ["agency", "region", "city", "title", "depart_date", "price", "discount_rate", "seats_left", "detail_url"]]
    top.columns = ["여행사", "지역", "도시", "상품명", "출발일", "판매가", "할인율", "잔여석", "링크"]

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame([{"항목": k, "값": v} for k, v in {
            "리포트 일자": summary["report_date"], "수집 상품 수": summary["total"], "여행사 수": summary["agencies"],
            "평균 판매가(원)": summary["avg_price"], "평균 할인율(%)": summary["avg_discount"], "최대 할인율(%)": summary["max_discount"],
        }.items()]).to_excel(w, index=False, sheet_name="요약")
        by_agency.to_excel(w, index=False, sheet_name="여행사별")
        by_region.to_excel(w, index=False, sheet_name="지역별")
        top.to_excel(w, index=False, sheet_name="할인TOP20")

        for ws in w.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0D6EFD")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(10, min(50, max(len(str(c.value or "")) for c in col) + 4))

        if len(df):
            ws = w.book.create_sheet("차트")
            fig, ax = plt.subplots(figsize=(8, 4))
            sns.countplot(data=df, x="region", hue="agency", ax=ax, order=by_region["지역"].tolist(), palette="Set2")
            ax.set(title="지역·여행사별 특가 상품 수", xlabel="지역", ylabel="상품 수")
            ws.add_image(XLImage(_png(fig)), "A1")
            fig, ax = plt.subplots(figsize=(8, 4))
            sns.boxplot(data=df, x="region", y="price", ax=ax, order=by_region["지역"].tolist(), color="#7FB3FF")
            ax.set(title="지역별 판매가 분포", xlabel="지역", ylabel="판매가(원)")
            ax.yaxis.set_major_formatter(ticker.FuncFormatter(lambda v, _: f"{v / 10000:,.0f}만"))
            ws.add_image(XLImage(_png(fig)), "A24")

    html = ""
    if len(df):
        f1 = px.bar(by_region, x="지역", y="상품수", color="평균할인율", title=f"{report_date} 지역별 특가 상품 수", color_continuous_scale="Blues")
        f2 = px.box(df, x="agency", y="price", color="agency", points="all", title="여행사별 판매가 분포",
                    labels={"agency": "여행사", "price": "판매가(원)"})
        html = ("<html><head><meta charset='utf-8'><title>NomadPick 일일 리포트</title></head><body>"
                + f1.to_html(full_html=False, include_plotlyjs="cdn") + f2.to_html(full_html=False, include_plotlyjs=False)
                + "</body></html>")
    return buf.getvalue(), html, summary
